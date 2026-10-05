# -*- coding: utf-8 -*-
"""HSBC 每日庫存損益報表 —— 核心引擎。

負責：讀 xlsx（唯讀）→ 抓收盤價 → 算市值/損益 → 用 ReportLab 畫出 PDF。
完全不需要 Microsoft Excel，也不需要 LibreOffice。

這個模組不碰任何 GUI，CLI 與 GUI 都呼叫 run_report()。
"""

from __future__ import annotations

import os
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

import openpyxl
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Flowable, SimpleDocTemplate, Table, TableStyle

REPORT_SHEET = "HSBC"
PRICE_SHEET = "工作表1"
DATE_CELL = "H4"
TOTAL_ROW_LABEL = "USD小計"

# 欄位（1-based）：D=單位數 G=申贖成本 H=收盤價 I=市值 J=P&L P=已實現損益
COL_UNITS, COL_COST = 4, 7
COL_PRICE, COL_MARKET_VALUE, COL_PNL = 8, 9, 10
COL_REALIZED_PNL = 16

PAGE_SIZE = A4
MARGIN = 22.7  # 約 0.31 吋，與原本 Excel 的頁面邊界一致


# ========================= 字型 =========================

FONT_LATIN = "Times-Roman"
FONT_LATIN_BOLD = "Times-Bold"
FONT_CJK = "Times-Roman"       # 會在 setup_fonts() 裡換掉
FONT_CJK_BOLD = "Times-Bold"

# 注意：ReportLab 只能嵌入「TrueType 外框」的字型。
# 像 Noto CJK 那種 .ttc/.otf（PostScript/CFF 外框）會載入失敗，不要放進來。
CJK_FONT_CANDIDATES = [
    (r"C:\Windows\Fonts\mingliu.ttc", 0),      # 細明體，原始表格用的就是 MingLiu
    (r"C:\Windows\Fonts\msjh.ttc", 0),         # 微軟正黑體
    (r"C:\Windows\Fonts\kaiu.ttf", None),      # 標楷體
    (r"C:\Windows\Fonts\simsun.ttc", 0),
    ("/usr/share/fonts/truetype/droid/DroidSansFallbackFull.ttf", None),
    ("/usr/share/fonts/truetype/arphic/uming.ttc", 0),
]


def setup_fonts(preferred: str | None = None):
    """註冊中文字型，回傳實際使用的字型路徑。

    一定要嵌入實體 TrueType 字型 —— ReportLab 內建的 CID 字型（MSung-Light）
    在很多 PDF 閱讀器上會整個顯示不出來（實測 poppler 直接空白），
    對一份每天自動產生、沒人盯著看的報表來說太危險，所以找不到就直接報錯。
    """
    global FONT_CJK, FONT_CJK_BOLD

    candidates = list(CJK_FONT_CANDIDATES)
    for override in (preferred, os.environ.get("HSBC_CJK_FONT")):
        if override:
            candidates.insert(0, (override, 0 if override.lower().endswith(".ttc") else None))

    tried = []
    for path, index in candidates:
        p = Path(path)
        if not p.exists():
            continue
        try:
            if index is None:
                pdfmetrics.registerFont(TTFont("CJK", str(p)))
            else:
                pdfmetrics.registerFont(TTFont("CJK", str(p), subfontIndex=index))
            if not pdfmetrics.getFont("CJK").face.charToGlyph.get(ord("購")):
                raise ValueError("這個字型沒有中文字（缺少「購」字）")
            FONT_CJK = FONT_CJK_BOLD = "CJK"
            return str(p)
        except Exception as e:
            tried.append(f"  {p}：{e}")

    detail = "\n".join(tried) if tried else "  （候選清單裡的字型檔都不存在）"
    raise RuntimeError(
        "找不到可用的中文字型，PDF 裡的中文會變成空白，因此停止執行。\n"
        f"已嘗試：\n{detail}\n"
        "請在「設定」頁指定一個 .ttf/.ttc 中文字型檔。"
    )


def split_runs(text):
    """把字串切成連續的「中文」與「非中文」區段，兩者分別用不同字型畫。

    有些中文字型（例如 Linux 的 DroidSansFallback）只收中文字、沒有英數字，
    整格直接套同一個字型的話，英文會整段消失不見。
    """
    runs = []
    for ch in text:
        is_cjk = ord(ch) > 0x2E80
        if runs and runs[-1][1] == is_cjk:
            runs[-1][0] += ch
        else:
            runs.append([ch, is_cjk])
    return runs


class MixedFontText(Flowable):
    """一格文字：中文與英數字各用各的字型，單行不換行。

    跟 Excel 一樣，文字太長就往右邊的空白格溢出（而不是折行）；
    max_width 是可溢出的總寬度，連溢出都放不下時就縮小字級，避免壓到旁邊有內容的格子。
    """

    def __init__(self, text, size, color, align, width, height, bold, max_width=None):
        super().__init__()
        self.text = text
        self.size = size
        self.color = color
        self.align = align
        self.width = width
        self.height = height
        self.bold = bold
        self.max_width = max_width

    def _runs(self):
        latin = FONT_LATIN_BOLD if self.bold else FONT_LATIN
        cjk = FONT_CJK_BOLD if self.bold else FONT_CJK
        return [(t, cjk if is_cjk else latin) for t, is_cjk in split_runs(self.text)]

    def wrap(self, available_width, available_height):
        return self.width, self.height

    def draw(self):
        runs = self._runs()
        size = self.size
        total = sum(pdfmetrics.stringWidth(t, f, size) for t, f in runs)
        if self.max_width and total > self.max_width:
            size *= self.max_width / total
            total = self.max_width

        if self.align == "RIGHT":
            x = self.width - total
        elif self.align == "CENTER":
            x = (self.width - total) / 2
        else:
            x = 0

        y = (self.height - size) / 2 + size * 0.22
        self.canv.setFillColor(self.color)
        for text, font in runs:
            self.canv.setFont(font, size)
            self.canv.drawString(x, y, text)
            x += pdfmetrics.stringWidth(text, font, size)


# ====================== 抓行情 ======================

def get_latest_closes(tickers, log=print):
    """逐檔下載日線，再找出所有股票都有收盤價的最新共同交易日。"""
    import yfinance as yf

    histories = {}
    for ticker in tickers:
        hist = yf.Ticker(ticker).history(period="1mo", interval="1d", auto_adjust=False)
        if hist.empty:
            raise RuntimeError(f"無法取得 {ticker} 行情")

        # yfinance 的 index 可能帶時區，統一只留日期；NaN（停牌日）直接跳過
        closes = {}
        for timestamp, close in hist["Close"].items():
            if close is None or close != close:
                continue
            closes[timestamp.date()] = float(close)

        if not closes:
            raise RuntimeError(f"{ticker} 近一個月沒有任何收盤價")

        histories[ticker] = closes
        log(f"{ticker} 最新可用日期：{max(closes)}")

    common_dates = set.intersection(*(set(c) for c in histories.values()))
    if not common_dates:
        raise RuntimeError("找不到所有庫存股票都有收盤價的共同交易日")

    trade_date = max(common_dates)
    log(f"最新共同美股收盤日：{trade_date:%Y-%m-%d}")

    prices = {t: histories[t][trade_date] for t in tickers}
    for ticker, price in prices.items():
        log(f"{ticker}: {trade_date:%Y-%m-%d} close = {price:.2f}")

    return prices, trade_date


# ====================== 讀 xlsx ======================

HOLDING_FORMULA = re.compile(rf"^={re.escape(PRICE_SHEET)}!B(\d+)$")


def find_holdings(ws_formulas, price_ws):
    """找出持股列：H 欄是 `=工作表1!Bn` 公式的列，代號取自工作表1 的 A 欄。"""
    holdings = {}
    for row in range(1, ws_formulas.max_row + 1):
        value = ws_formulas.cell(row=row, column=COL_PRICE).value
        if not isinstance(value, str):
            continue
        match = HOLDING_FORMULA.match(value.strip())
        if not match:
            continue
        ticker = price_ws.cell(row=int(match.group(1)), column=1).value
        if ticker:
            holdings[row] = str(ticker).strip()

    if not holdings:
        raise RuntimeError(
            f"在 {REPORT_SHEET} 工作表的 H 欄找不到任何 `={PRICE_SHEET}!Bn` 公式，"
            "無法判斷目前持股有哪些"
        )
    return holdings


def find_total_row(ws_values):
    """找最底下的小計列（B 欄標籤）。"""
    for row in range(1, ws_values.max_row + 1):
        if ws_values.cell(row=row, column=2).value == TOTAL_ROW_LABEL:
            return row
    return None


def build_values(ws_values, holdings, prices, trade_date, log=print):
    """把 xlsx 的快取值複製一份，再用最新行情覆蓋會變動的那些格子。"""
    max_row, max_col = ws_values.max_row, ws_values.max_column
    grid = {
        (r, c): ws_values.cell(row=r, column=c).value
        for r in range(1, max_row + 1)
        for c in range(1, max_col + 1)
    }

    grid[(ws_values[DATE_CELL].row, ws_values[DATE_CELL].column)] = trade_date

    rows = []
    total_pnl = 0.0
    for row, ticker in sorted(holdings.items()):
        price = prices[ticker]
        units = grid.get((row, COL_UNITS))
        cost = grid.get((row, COL_COST))
        if not isinstance(units, (int, float)) or not isinstance(cost, (int, float)):
            raise RuntimeError(
                f"{REPORT_SHEET}!列{row}（{ticker}）的單位數或申贖成本不是數字，無法計算市值。\n"
                "請先用 Excel 開啟這個檔案存一次，讓公式算出結果。"
            )

        market_value = price * units
        pnl = market_value - cost
        total_pnl += pnl

        grid[(row, COL_PRICE)] = price
        grid[(row, COL_MARKET_VALUE)] = market_value
        grid[(row, COL_PNL)] = pnl
        rows.append(
            {
                "ticker": ticker,
                "price": price,
                "units": units,
                "cost": cost,
                "market_value": market_value,
                "pnl": pnl,
            }
        )
        log(f"{ticker} (列{row}): 收盤價={price:.2f} 單位數={units:g} "
            f"市值={market_value:,.2f} P&L={pnl:,.2f}")

    realized = 0.0
    total_row = find_total_row(ws_values)
    if total_row:
        grid[(total_row, COL_PNL)] = total_pnl
        realized = sum(
            v for r in range(1, total_row)
            if isinstance(v := grid.get((r, COL_REALIZED_PNL)), (int, float))
        )
        grid[(total_row, COL_REALIZED_PNL)] = realized
        log(f"庫存損益小計={total_pnl:,.2f}  已實現損益小計={realized:,.2f}")

    return grid, max_row, max_col, rows, total_pnl, realized


# ==================== 數字/日期格式 ====================

def format_value(value, number_format):
    """照 Excel 的數字格式字串把值轉成顯示文字（只處理這張表會用到的格式）。"""
    if value is None:
        return ""

    if isinstance(value, (datetime, date)):
        # openpyxl 把 Excel 內建格式 14 報成 'mm-dd-yy'，但那其實是「系統短日期」，
        # 在繁體中文 Windows 的 Excel 裡顯示的是 yyyy/mm/dd，所以一律照後者輸出。
        return value.strftime("%Y/%m/%d")

    if not isinstance(value, (int, float)):
        return str(value)

    fmt = number_format.split(";")[0] if number_format else "General"

    if "%" in number_format:
        return f"{value * 100:.2f}%"

    negative_in_parens = "(" in number_format
    text_value = abs(value) if negative_in_parens and value < 0 else value

    if "#,##0.00" in fmt:
        text = f"{text_value:,.2f}"
    elif "#,##0" in fmt:
        text = f"{text_value:,.0f}"
    elif fmt == "0.00":
        text = f"{text_value:.2f}"
    elif fmt == "0":
        text = f"{text_value:.0f}"
    else:  # General
        text = f"{text_value:g}" if isinstance(text_value, float) else str(text_value)

    if negative_in_parens and value < 0:
        return f"({text})"
    return text


def is_red_negative(value, number_format):
    """Excel 格式字串裡標了 [Red] 的負數要用紅字。"""
    return isinstance(value, (int, float)) and value < 0 and "[Red]" in number_format


# ====================== 樣式轉換 ======================

def argb_to_color(argb):
    """openpyxl 的 'FFRRGGBB' 轉成 ReportLab 顏色；主題色/不確定的就回 None。"""
    if not isinstance(argb, str) or len(argb) != 8:
        return None
    try:
        return colors.HexColor("#" + argb[2:])
    except ValueError:
        return None


def cell_font_color(cell):
    return argb_to_color(getattr(cell.font.color, "rgb", None)) or colors.black


def cell_background(cell):
    fill = cell.fill
    if fill is None or fill.fill_type != "solid":
        return None
    color = argb_to_color(getattr(fill.fgColor, "rgb", None))
    if color is None or color == colors.white:
        return None
    return color


ALIGN_MAP = {"left": "LEFT", "right": "RIGHT", "center": "CENTER"}


def default_align(value):
    """Excel 沒指定對齊時：數字/日期靠右，文字靠左。"""
    if isinstance(value, (int, float, datetime, date)):
        return "RIGHT"
    return "LEFT"


# ======================== 畫 PDF ========================

def excel_width_to_points(width):
    """Excel 欄寬（字元數）換算成點：像素 = 字元數×7 + 5，再 ×0.75 轉點。"""
    return (width * 7 + 5) * 0.75


def render_pdf(ws, grid, max_row, max_col, trade_date, pdf_path):
    """用 ReportLab 把整張表畫成 PDF，樣式沿用 xlsx 裡的設定。"""
    default_width = ws.sheet_format.defaultColWidth or 8.43
    col_widths = []
    for c in range(1, max_col + 1):
        letter = openpyxl.utils.get_column_letter(c)
        dim = ws.column_dimensions.get(letter)
        width = dim.width if dim and dim.width else default_width
        col_widths.append(excel_width_to_points(width))

    # 整張表縮放到剛好塞進頁寬（原本 Excel 頁面設定是縮放 57%，這裡自動算）
    usable_width = PAGE_SIZE[0] - 2 * MARGIN
    scale = min(1.0, usable_width / sum(col_widths))
    col_widths = [w * scale for w in col_widths]

    # 合併儲存格要用整個合併範圍的寬度來對齊，不然置中的標題會黏在最左欄
    merged_width = {}
    for merged in ws.merged_cells.ranges:
        if merged.max_col <= max_col:
            merged_width[(merged.min_row, merged.min_col)] = sum(
                col_widths[merged.min_col - 1 : merged.max_col]
            )

    default_row_height = ws.sheet_format.defaultRowHeight or 13.2
    row_heights = []
    data = []
    style = [
        ("TOPPADDING", (0, 0), (-1, -1), 0),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
        ("LEFTPADDING", (0, 0), (-1, -1), 1.5 * scale),
        ("RIGHTPADDING", (0, 0), (-1, -1), 1.5 * scale),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]

    for r in range(1, max_row + 1):
        dim = ws.row_dimensions.get(r)
        height = dim.height if dim and dim.height else default_row_height
        row_heights.append(height * scale)

        row_data = []
        for c in range(1, max_col + 1):
            cell = ws.cell(row=r, column=c)
            value = grid.get((r, c))
            number_format = cell.number_format or "General"
            text = format_value(value, number_format)

            pos = (c - 1, r - 1)
            bold = bool(cell.font.bold)
            font_size = (cell.font.sz or 10) * scale
            color = colors.red if is_red_negative(value, number_format) else cell_font_color(cell)
            align = ALIGN_MAP.get(cell.alignment.horizontal) or default_align(value)

            if any(ord(ch) > 0x2E80 for ch in text):
                # 含中文：交給 MixedFontText 自己畫，中英文各用各的字型
                cell_width = merged_width.get((r, c), col_widths[c - 1])
                max_width = None
                if align == "LEFT":
                    # 跟 Excel 一樣只能溢出到右邊連續的空白格，碰到有內容的格子就停
                    overflow = cell_width
                    for nc in range(c + 1, max_col + 1):
                        if grid.get((r, nc)) not in (None, ""):
                            break
                        overflow += col_widths[nc - 1]
                    max_width = overflow - 3 * scale
                row_data.append(
                    MixedFontText(
                        text, font_size, color, align,
                        cell_width - 3 * scale, height * scale, bold, max_width,
                    )
                )
            else:
                row_data.append(text)
                style.append(("FONT", pos, pos, FONT_LATIN_BOLD if bold else FONT_LATIN, font_size))
                style.append(("TEXTCOLOR", pos, pos, color))
                style.append(("ALIGN", pos, pos, align))

            background = cell_background(cell)
            if background is not None:
                style.append(("BACKGROUND", pos, pos, background))

            border = cell.border
            for side, command in (
                ("left", "LINEBEFORE"), ("right", "LINEAFTER"),
                ("top", "LINEABOVE"), ("bottom", "LINEBELOW"),
            ):
                side_border = getattr(border, side)
                if not side_border or not side_border.style:
                    continue
                weight = 1.2 if side_border.style in ("medium", "thick") else 0.5
                line_color = argb_to_color(getattr(side_border.color, "rgb", None)) or colors.black
                style.append((command, pos, pos, weight * scale, line_color))

        data.append(row_data)

    for merged in ws.merged_cells.ranges:
        if merged.max_row > max_row or merged.max_col > max_col:
            continue
        style.append(
            ("SPAN",
             (merged.min_col - 1, merged.min_row - 1),
             (merged.max_col - 1, merged.max_row - 1))
        )

    doc = SimpleDocTemplate(
        str(pdf_path),
        pagesize=PAGE_SIZE,
        leftMargin=MARGIN, rightMargin=MARGIN,
        topMargin=MARGIN * 1.6, bottomMargin=MARGIN * 1.6,
        title=f"HSBC List {trade_date:%Y%m%d}",
        author="HSBC Daily Report",
    )
    table = Table(data, colWidths=col_widths, rowHeights=row_heights, repeatRows=0)
    table.setStyle(TableStyle(style))
    doc.build([table])
    return pdf_path


# ======================== 對外介面 ========================

@dataclass
class ReportResult:
    trade_date: date
    pdf_path: Path
    holdings: list = field(default_factory=list)
    total_pnl: float = 0.0
    realized_pnl: float = 0.0


def run_report(excel_path: Path, pdf_dir: Path, cjk_font: str | None = None, log=print) -> ReportResult:
    """跑完整流程：讀 xlsx → 抓行情 → 算損益 → 產生 PDF。

    xlsx 全程唯讀，絕對不會寫回去。
    """
    excel_path = Path(excel_path)
    pdf_dir = Path(pdf_dir)

    if not excel_path.exists():
        raise FileNotFoundError(f"找不到資料來源 Excel：{excel_path}")
    pdf_dir.mkdir(parents=True, exist_ok=True)

    log(f"中文字型：{setup_fonts(cjk_font)}")

    wb_formulas = openpyxl.load_workbook(excel_path, data_only=False)
    wb_values = openpyxl.load_workbook(excel_path, data_only=True)
    ws_formulas = wb_formulas[REPORT_SHEET]
    ws_values = wb_values[REPORT_SHEET]

    holdings = find_holdings(ws_formulas, wb_formulas[PRICE_SHEET])
    log("目前持股：" + "、".join(f"{t}(列{r})" for r, t in sorted(holdings.items())))

    tickers = list(dict.fromkeys(holdings.values()))
    prices, trade_date = get_latest_closes(tickers, log=log)

    grid, max_row, max_col, rows, total_pnl, realized = build_values(
        ws_values, holdings, prices, trade_date, log=log
    )

    pdf_path = pdf_dir / f"HSBC List_{trade_date:%Y%m%d}.pdf"
    if pdf_path.exists():
        try:
            pdf_path.unlink()
        except PermissionError:
            raise RuntimeError(f"PDF 正在被其他程式開啟，請先關閉後再執行：\n{pdf_path}")

    render_pdf(ws_values, grid, max_row, max_col, trade_date, pdf_path)
    log(f"PDF 已建立：{pdf_path}")

    return ReportResult(
        trade_date=trade_date,
        pdf_path=pdf_path,
        holdings=rows,
        total_pnl=total_pnl,
        realized_pnl=realized,
    )
